import re, sys, urllib.parse
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

src = sys.argv[1]
txt = open(src, encoding='utf-8').read()
# direcciones tal como las da Google (no adivinar formato)
txt = txt.replace('33 Avenida 27-40', '33 Avenida 2740').replace('23 Av 28-26', '23 Av 2826').replace('0 Avenida 2-35', '0 Avenida 235')
rows = [l.split('|') for l in txt.splitlines() if l.strip()]

MSG = ("Buenas, le saluda Javier Arriaga, distribuidor de camarón congelado. Manejo PYD, cola y entero en varias tallas, "
       "peso neto, y entrego en su negocio con pago contra entrega. ¿Cuántas libras de camarón trabajan por semana y qué talla usan? "
       "Con eso le paso precio.")
COMP = {'Blue Ocean Seafood Guatemala', 'Gambore, S.A.', 'Mariscos, Pescados y Carnes Frescas', 'Lo Del Mar', 'MAR Y TIERRA PURPLE, S.A.'}
REV = {'Mariscos Damián', 'Mariscos El Robalo', 'Mar Mad', 'Mariscos De León SAN CRISTÓBAL', 'Productos Frescos del Mar'}


def area(lat, lng, d):
    d = d.lower()
    if 'amatit' in d or (lng < -90.6 and lat < 14.5):
        return 'Amatitlán'
    if 'villa nueva' in d:
        return 'Villa Nueva (límite Petapa)'
    if lat < 14.535:
        return 'San Miguel Petapa / Villa Nueva'
    if lat < 14.56:
        return 'Villa Nueva / Av. Petapa'
    if 'mixco' in d or lng < -90.58:
        return 'Mixco'
    return 'Ciudad capital'


out, comp = [], []
for q, name, cat, dirr, tel, web, rev, ll in rows:
    lat, lng = map(float, ll.split(','))
    m = re.search(r'\((\d+)\)', rev)
    n = int(m.group(1)) if m else 0
    est = rev.split('(')[0] if rev else ''
    ar = area(lat, lng, dirr)
    mapa = f'https://maps.google.com/?q={ll}'
    if name in COMP:
        comp.append([name, cat, dirr.strip(), ar, tel, web, mapa])
        continue
    tipo = 'Revendedor (pescadería/mercado)' if name in REV else 'Restaurante / cevichería'
    cadena = any(k in name.lower() for k in ['chavos', 'chatos', 'chatio', 'chatito', 'muelle viejo'])
    if n >= 300 or (cadena and n >= 50):
        pri = 'A'
    elif name in REV or n >= 50:
        pri = 'B'
    else:
        pri = 'C'
    movil = bool(tel) and tel[0] in '345'
    wa = f'https://wa.me/502{tel}?text=' + urllib.parse.quote(MSG) if movil else ''
    canal = 'WhatsApp' if movil else ('Llamar (fijo)' if tel else 'Visitar / Facebook')
    zc = ''
    if ar == 'Ciudad capital':
        mz = re.search(r'[Zz]ona (\d+)', dirr)
        zc = 'Zona ' + mz.group(1) if mz else (q.title() + ' (aprox.)' if q.startswith('zona') else '')
    nota = []
    if cadena:
        nota.append('Cadena con varias sucursales: pedir contacto del dueño/compras, un solo trato para todas')
    if name in REV:
        nota.append('Compra para revender: pedido por volumen, rota rápido')
    if n >= 1000:
        nota.append('Muy alto movimiento (+1,000 reseñas): probable consumo alto')
    if not tel:
        nota.append('Sin teléfono en Google: ir en persona o buscar en Facebook')
    out.append([pri, name, tipo, ar, zc, dirr.strip(), 'No aparece (preguntar en visita)',
                f'{tel[:4]}-{tel[4:]}' if tel else '', canal, wa, web, mapa, est, n, '; '.join(nota),
                'Por contactar', '', '', '', '', '', ''])

out.sort(key=lambda r: ({'A': 0, 'B': 1, 'C': 2}[r[0]], r[3], -r[13]))

wb = Workbook()
ws = wb.active
ws.title = 'Clientes'
H = ['Prioridad', 'Negocio', 'Tipo', 'Área', 'Zona', 'Dirección (Google Maps)', 'Propietario', 'Teléfono', 'Cómo contactar',
     'WhatsApp con mensaje listo', 'Web / Facebook', 'Ubicación exacta (mapa)', 'Estrellas', 'Reseñas (movimiento)', 'Notas',
     'Estado', 'Fecha contacto', 'Libras/semana', 'Producto y talla que usa', 'Precio que paga hoy', 'Proveedor actual', 'Próximo paso']
ws.append(H)
hf = PatternFill('solid', fgColor='1F4E78')
gris = PatternFill('solid', fgColor='EDEDED')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = hf
    c.alignment = Alignment(wrap_text=True, vertical='center')
fills = {'A': 'C6EFCE', 'B': 'FFEB9C', 'C': 'F2F2F2'}
for r in out:
    ws.append(r)
    i = ws.max_row
    ws.cell(i, 1).fill = PatternFill('solid', fgColor=fills[r[0]])
    ws.cell(i, 1).font = Font(bold=True)
    for col, label in ((10, 'Abrir WhatsApp'), (11, None), (12, 'Ver en mapa')):
        v = ws.cell(i, col).value
        if v:
            ws.cell(i, col).hyperlink = v if v.startswith('http') else 'https://' + v.split(' ')[0]
            ws.cell(i, col).value = label or v
            ws.cell(i, col).font = Font(color='0563C1', underline='single')
    for col in range(16, 23):
        ws.cell(i, col).fill = gris
for k, w in enumerate([9, 34, 22, 24, 15, 42, 22, 12, 16, 16, 30, 13, 9, 11, 50, 14, 13, 12, 22, 15, 18, 22], 1):
    ws.column_dimensions[ws.cell(1, k).column_letter].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type='list', formula1='"Por contactar,Contactado,Respondió,Muestra entregada,Cliente,No interesa"')
ws.add_data_validation(dv)
dv.add(f'P2:P{ws.max_row}')

c2 = wb.create_sheet('Competencia-Proveedores')
c2.append(['Negocio', 'Tipo', 'Dirección', 'Área', 'Teléfono', 'Web', 'Mapa', 'Para qué sirve'])
for r in comp:
    c2.append(r + ['NO es cliente: distribuidor. Sirve para comparar precios' + (' (tu proveedor actual)' if 'Blue Ocean' in r[0] else '')])
for c in c2[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = hf
for k, w in enumerate([34, 22, 45, 18, 12, 32, 40, 60], 1):
    c2.column_dimensions[c2.cell(1, k).column_letter].width = w

g = wb.create_sheet('Cómo usar')
for line in [
    'CÓMO TRABAJAR ESTA LISTA',
    '1. Empieza por los A (verde): más reseñas = más clientela = más libras. Luego los B (amarillo).',
    '2. Celular → clic en "Abrir WhatsApp" (mensaje ya escrito, solo das Enviar). Fijo → llamar. Sin teléfono → visitar.',
    '3. Escribir después de las 8 am; mejor 9-11 am o 3-5 pm (fuera de la hora de almuerzo).',
    '4. Primer contacto SIN precio. Primero saber: libras por semana, talla, producto, proveedor actual y cuánto paga. Anótalo en las columnas grises.',
    '5. Califica: 25+ lb/semana = cliente A real. Menos de 10 lb = menudeo, no vale el viaje.',
    '6. Entrada: llevar muestra de 2 lb de la talla que usa; primer pedido con pago contra entrega.',
    '7. Visitas por ruta para ahorrar gasolina: filtra la columna Área (Petapa/Villa Nueva un día, zonas 5-18 otro, zonas 11-12 otro, Amatitlán otro).',
    '8. Cadenas (Los Chavos, Los Chatos/Chatios, Muelle Viejo): pide al encargado el contacto del dueño o de compras.',
    '9. Meta Q10,000/mes: 5-6 clientes de ~30 lb/semana. Trabajando 15-20 a fondo deberías cerrar esos 5-6.',
    '',
    'AVISOS',
    '• Datos sacados de Google Maps el 27 oct 2026. Confirmar antes de ir: pueden haber cerrado o cambiado de número.',
    '• El nombre del dueño NO sale en Google Maps: se pregunta en la visita ("¿con quién trato lo de proveedores?").',
    '• Algunas direcciones salen como código de Google (ej. JC2Q+4CW): el botón "Ver en mapa" lleva al punto exacto.',
    '• Zona con "(aprox.)" = sale de la búsqueda; la dirección no traía zona.',
]:
    g.append([line])
g['A1'].font = Font(bold=True, size=14)
g['A12'].font = Font(bold=True, size=12)
g.column_dimensions['A'].width = 125

p = sys.argv[2]
wb.save(p)
print(len(out), len(comp), Counter(r[0] for r in out), Counter(r[3] for r in out),
      'tel', sum(1 for r in out if r[7]), 'wa', sum(1 for r in out if r[9]))
